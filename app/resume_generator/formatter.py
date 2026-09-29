"""Render a TailoredResume to ATS-friendly PDF and DOCX.

ATS rules (both formats): single column, no tables/text boxes/headers/footers/
images, standard section headings in a fixed order, dates right-aligned on the
company/title line (tab stop in DOCX, a right-aligned text run in PDF — never
a table), skills as a comma-separated paragraph, real "•" bullets, Letter
size with 0.75" margins.
"""

from __future__ import annotations

import logging
from pathlib import Path
from xml.sax.saxutils import escape

from app.resume_generator.models import TailoredResume

logger = logging.getLogger(__name__)

SECTION_ORDER = ["Summary", "Skills", "Experience", "Education", "Certifications"]
_MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
               "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def format_month(value: str | None) -> str:
    """"2022-04" → "Apr 2022"; "2022" → "2022"; None/"" → ""."""
    if not value:
        return ""
    parts = value.split("-")
    if len(parts) >= 2 and parts[1].isdigit() and 1 <= int(parts[1]) <= 12:
        return f"{_MONTH_ABBR[int(parts[1]) - 1]} {parts[0]}"
    return parts[0]


def date_range(start: str | None, end: str | None, current_label: str = "Present") -> str:
    start_s = format_month(start)
    end_s = format_month(end) if end else current_label
    return f"{start_s} – {end_s}" if start_s else end_s


def _position_left(title: str, company: str, location: str | None) -> tuple[str, str]:
    """(bold part, regular part) for the title/company line."""
    rest = f", {company}" + (f" — {location}" if location else "")
    return title, rest


def _education_lines(resume: TailoredResume) -> list[tuple[str, str, str]]:
    lines = []
    for e in resume.education:
        detail = ", ".join(x for x in (e.degree, e.field_of_study) if x)
        years = " – ".join(y for y in (e.start_year, e.end_year) if y)
        lines.append((e.school, f", {detail}" if detail else "", years))
    return lines


def _cert_lines(resume: TailoredResume) -> list[tuple[str, str, str]]:
    return [
        (c.name, f", {c.authority}" if c.authority else "", format_month(c.issued))
        for c in resume.certifications
    ]


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------

def render_pdf(resume: TailoredResume, out_path: Path) -> Path:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import (
        Flowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer,
    )

    body_size = 10.5
    body = ParagraphStyle("Body", fontName="Helvetica", fontSize=body_size, leading=13.5)
    name_style = ParagraphStyle("Name", fontName="Helvetica-Bold", fontSize=15, leading=19,
                                alignment=TA_CENTER)
    center = ParagraphStyle("Center", parent=body, alignment=TA_CENTER)
    heading = ParagraphStyle("Heading", fontName="Helvetica-Bold", fontSize=11.5, leading=15,
                             spaceBefore=9, spaceAfter=3)
    bullet_style = ParagraphStyle("Bullet", parent=body, leading=13)

    class LeftRightLine(Flowable):
        """Bold+regular text on the left, right-aligned text on the same line."""

        def __init__(self, bold: str, regular: str, right: str) -> None:
            super().__init__()
            self.bold, self.regular, self.right = bold, regular, right
            self.height = body.leading

        def wrap(self, avail_width: float, avail_height: float) -> tuple[float, float]:
            self.width = avail_width
            return avail_width, self.height

        def draw(self) -> None:
            c = self.canv
            y = self.height - body_size
            bold_w = stringWidth(self.bold, "Helvetica-Bold", body_size)
            c.setFont("Helvetica-Bold", body_size)
            c.drawString(0, y, self.bold)
            c.setFont("Helvetica", body_size)
            c.drawString(bold_w, y, self.regular)
            c.drawRightString(self.width, y, self.right)

    def entry_line(bold: str, regular: str, right: str, width: float) -> list:
        """One line if it fits; otherwise a wrapped paragraph with the date on its own line."""
        needed = (stringWidth(bold, "Helvetica-Bold", body_size)
                  + stringWidth(regular + "    " + right, "Helvetica", body_size))
        if needed <= width:
            return [LeftRightLine(bold, regular, right)]
        return [Paragraph(f"<b>{escape(bold)}</b>{escape(regular)}", body),
                Paragraph(escape(right), ParagraphStyle("Right", parent=body, alignment=2))]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    margin = 0.75 * inch
    doc = SimpleDocTemplate(
        str(out_path), pagesize=letter,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
        title=f"{resume.full_name} — Resume", author=resume.full_name,
    )
    width = letter[0] - 2 * margin

    story: list = [Paragraph(escape(resume.full_name), name_style)]
    if resume.headline:
        story.append(Paragraph(escape(resume.headline), center))
    if resume.contact_line:
        story.append(Paragraph(escape(resume.contact_line), center))

    if resume.summary:
        story += [Paragraph("Summary", heading), Paragraph(escape(resume.summary), body)]

    if resume.skills:
        story += [Paragraph("Skills", heading), Paragraph(escape(", ".join(resume.skills)), body)]

    if resume.positions:
        story.append(Paragraph("Experience", heading))
        for i, p in enumerate(resume.positions):
            if i:
                story.append(Spacer(1, 5))
            bold, regular = _position_left(p.title, p.company, p.location)
            story += entry_line(bold, regular, date_range(p.start_date, p.end_date), width)
            if p.bullets:
                story.append(ListFlowable(
                    [ListItem(Paragraph(escape(b.text), bullet_style), leftIndent=12)
                     for b in p.bullets],
                    bulletType="bullet", start="•", leftIndent=12, bulletFontSize=body_size,
                ))

    if resume.education:
        story.append(Paragraph("Education", heading))
        for bold, regular, right in _education_lines(resume):
            story += entry_line(bold, regular, right, width)

    if resume.certifications:
        story.append(Paragraph("Certifications", heading))
        for bold, regular, right in _cert_lines(resume):
            story += entry_line(bold, regular, right, width)

    doc.build(story)
    logger.info("Rendered resume PDF → %s", out_path)
    return out_path


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def render_docx(resume: TailoredResume, out_path: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.75))
    text_width = Inches(8.5 - 1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    # East-Asian font slot too, or Word may substitute
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
                   "Calibri")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    def para(text: str = "", bold: bool = False, size: float | None = None,
             center: bool = False, space_before: float = 0, space_after: float = 0):
        p = doc.add_paragraph()
        if text:
            run = p.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def heading(text: str) -> None:
        para(text, bold=True, size=11.5, space_before=9, space_after=3)

    def entry_line(bold: str, regular: str, right: str) -> None:
        p = para()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(bold).bold = True
        p.add_run(regular)
        if right:
            p.add_run("\t" + right)

    para(resume.full_name, bold=True, size=15, center=True)
    if resume.headline:
        para(resume.headline, center=True)
    if resume.contact_line:
        para(resume.contact_line, center=True)

    if resume.summary:
        heading("Summary")
        para(resume.summary)

    if resume.skills:
        heading("Skills")
        para(", ".join(resume.skills))

    if resume.positions:
        heading("Experience")
        for i, p in enumerate(resume.positions):
            bold, regular = _position_left(p.title, p.company, p.location)
            entry_line(bold, regular, date_range(p.start_date, p.end_date))
            if i:
                doc.paragraphs[-1].paragraph_format.space_before = Pt(5)
            for b in p.bullets:
                bp = para("•\t" + b.text)
                fmt = bp.paragraph_format
                fmt.left_indent = Inches(0.25)
                fmt.first_line_indent = Inches(-0.15)
                fmt.tab_stops.add_tab_stop(Inches(0.25))

    if resume.education:
        heading("Education")
        for bold, regular, right in _education_lines(resume):
            entry_line(bold, regular, right)

    if resume.certifications:
        heading("Certifications")
        for bold, regular, right in _cert_lines(resume):
            entry_line(bold, regular, right)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    logger.info("Rendered resume DOCX → %s", out_path)
    return out_path
