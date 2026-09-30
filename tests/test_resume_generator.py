"""Tests for the LinkedIn parser, anti-fabrication check, formatter, and resume pipeline.

Run: pytest tests/test_resume_generator.py
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

from app.config import (
    AppSettings, BoardsConfig, ContactConfig, ProfileConfig, SearchLaneConfig, SecretsConfig,
)
from app.database import Database, JobRecord
from app.polisher.claude_client import ResumeGenResult
from app.resume_generator.builder import assemble_resume, check_fabrication
from app.resume_generator.formatter import render_docx, render_pdf
from app.resume_generator.linkedin_parser import (
    load_additional_bullets,
    load_linkedin_data,
    merge_additional_bullets,
    normalize_date,
    parse_linkedin_export,
    split_bullets,
)
from app.resume_generator.models import ResumeBullet, ResumePosition, TailoredResume


# ---------------------------------------------------------------------------
# Fixtures — headers match a real LinkedIn export (Sep 2026)
# ---------------------------------------------------------------------------

def _csv(header: list[str], rows: list[list[str]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    return "﻿" + buf.getvalue()  # LinkedIn's UTF-8 BOM


@pytest.fixture
def export_zip(tmp_path: Path) -> Path:
    path = tmp_path / "linkedin_export.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("Profile.csv", _csv(
            ["First Name", "Last Name", "Maiden Name", "Address", "Birth Date", "Headline",
             "Summary", "Industry", "Zip Code", "Geo Location", "Twitter Handles", "Websites",
             "Instant Messengers"],
            [["Ada", "Lovelace", "", "", "", "Frontend Lead", "Builds things.", "IT", "60430",
              "Greater Chicago Area", "", "[PORTFOLIO:https://ada.dev]", ""]],
        ))
        zf.writestr("Positions.csv", _csv(
            ["Company Name", "Title", "Description", "Location", "Started On", "Finished On"],
            [
                ["Old Co", "Web Developer", "• Built pages in HTML and CSS • Wrote jQuery",
                 "Chicago, IL", "2015", "Mar 2016"],
                ["Acme Corp", "Senior Frontend Developer  ",
                 "•\tLed React migration for 12 apps •\tManaged day-to-day releases",
                 "Remote", "Apr 2022", ""],
                ["Mid Inc", "Developer", "Line one\n- Line two with AEM\n* Line three",
                 "", "Jun 2018", "Dec 2021"],
            ],
        ))
        zf.writestr("Skills.csv", _csv(["Name"], [["JavaScript"], ["React"], ["AEM"], ["React"]]))
        zf.writestr("Education.csv", _csv(
            ["School Name", "Start Date", "End Date", "Notes", "Degree Name", "Activities"],
            [["State University", "1995", "2000", "", "BS", ""]],
        ))
        zf.writestr("Certifications.csv", _csv(
            ["Name", "Url", "Authority", "Started On", "Finished On", "License Number"],
            [["AWS Certified AI Practitioner", "", "Amazon Web Services (AWS)", "Oct 2025", "",
              "x"]],
        ))
        zf.writestr("Email Addresses.csv", _csv(
            ["Email Address", "Confirmed", "Primary", "Updated On"],
            [["other@example.com", "Yes", "No", ""], ["ada@example.com", "Yes", "Yes", ""]],
        ))
        zf.writestr("PhoneNumbers.csv", _csv(["Extension", "Number", "Type"],
                                             [["", "555-0100", "Mobile"]]))
        zf.writestr("Connections.csv", _csv(["First Name"], [["Ignored"]]))
    return path


@pytest.fixture
def linkedin_data(export_zip: Path, tmp_path: Path):
    return parse_linkedin_export(
        export_zip, contact=ContactConfig(linkedin_url="https://linkedin.com/in/ada"),
        out_path=tmp_path / "linkedin_data.json",
    )


# ---------------------------------------------------------------------------
# Parser
# ---------------------------------------------------------------------------

def test_normalize_date() -> None:
    assert normalize_date("Apr 2022") == "2022-04"
    assert normalize_date("April 2022") == "2022-04"
    assert normalize_date("2015") == "2015-01"
    assert normalize_date("") is None
    with pytest.raises(ValueError):
        normalize_date("sometime")


def test_split_bullets_inline_and_multiline() -> None:
    assert split_bullets("•\tOne •\tTwo day-to-day") == ["One", "Two day-to-day"]
    assert split_bullets("A\n- B\n* C") == ["A", "B", "C"]
    assert split_bullets("") == []


def test_parser_builds_linkedin_data(linkedin_data, tmp_path: Path) -> None:
    d = linkedin_data
    assert d.full_name == "Ada Lovelace"
    assert d.headline == "Frontend Lead"
    assert d.location == "Greater Chicago Area"
    assert d.email == "ada@example.com"          # primary address wins
    assert d.phone == "555-0100"
    assert d.linkedin_url == "https://linkedin.com/in/ada"
    assert d.websites == ["https://ada.dev"]
    assert d.skills == ["JavaScript", "React", "AEM"]  # de-duplicated
    assert d.all_skills_lower == {"javascript", "react", "aem"}

    # Newest first, dates normalized, current role has no end date
    assert [p.company for p in d.positions] == ["Acme Corp", "Mid Inc", "Old Co"]
    acme, _, old = d.positions
    assert acme.title == "Senior Frontend Developer"  # whitespace stripped
    assert (acme.start_date, acme.end_date) == ("2022-04", None)
    assert (old.start_date, old.end_date) == ("2015-01", "2016-03")
    assert acme.bullets == ["Led React migration for 12 apps", "Managed day-to-day releases"]

    assert d.education[0].school == "State University"
    assert (d.education[0].start_year, d.education[0].end_year) == ("1995", "2000")
    assert d.certifications[0].issued == "2025-10"
    assert d.parse_warnings == []

    # Written to disk and loadable
    reloaded = load_linkedin_data(tmp_path / "linkedin_data.json")
    assert reloaded.positions == d.positions


def test_parser_warns_on_missing_contact(export_zip: Path) -> None:
    d = parse_linkedin_export(export_zip, contact=ContactConfig(), out_path=None)
    assert any("linkedin_url" in w for w in d.parse_warnings)


def test_load_linkedin_data_missing(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="parse-linkedin"):
        load_linkedin_data(tmp_path / "nope.json")


# ---------------------------------------------------------------------------
# Additional bullets
# ---------------------------------------------------------------------------

def test_additional_bullets_merge(linkedin_data, tmp_path: Path) -> None:
    path = tmp_path / "additional_bullets.json"
    path.write_text(json.dumps({
        "_readme": "ignored",
        "_example": {"Acme Corp": [{"text": "Should never merge", "tags": []}]},
        "acme corp": [{"text": "Cut build time 40% with Vite", "tags": ["Vite", "Performance"]}],
    }))
    bullets = load_additional_bullets(path)
    assert list(bullets) == ["acme corp"]

    merged = merge_additional_bullets(linkedin_data, bullets)
    acme = merged.positions[0]
    assert acme.bullets[-1] == "Cut build time 40% with Vite"
    assert "Should never merge" not in acme.bullets
    assert merged.positions[1].bullets == linkedin_data.positions[1].bullets  # untouched
    assert merged.supplemental_tags == ["Vite", "Performance"]
    # Original is not mutated
    assert "Cut build time 40% with Vite" not in linkedin_data.positions[0].bullets


def test_additional_bullets_missing_file(tmp_path: Path) -> None:
    assert load_additional_bullets(tmp_path / "missing.json") == {}


# ---------------------------------------------------------------------------
# Anti-fabrication
# ---------------------------------------------------------------------------

def _resume(**overrides) -> TailoredResume:
    base = dict(
        full_name="Ada Lovelace", headline="h", contact_line="c", summary="s",
        positions=[ResumePosition(
            company="Acme Corp", title="Senior Frontend Developer", start_date="2022-04",
            bullets=[ResumeBullet(text="Led React migration", matched_skills=["React", "Rust"])],
        )],
        skills=["React", "JavaScript", "Kubernetes", "jQuery"],
    )
    base.update(overrides)
    return TailoredResume(**base)


def test_fabricated_skill_removed(linkedin_data) -> None:
    resume = check_fabrication(_resume(), linkedin_data)
    assert "Kubernetes" not in resume.skills
    assert "jQuery" in resume.skills  # found in a position bullet, not the skills list
    assert resume.positions[0].bullets[0].matched_skills == ["React"]
    assert any("Kubernetes" in w for w in resume.fabrication_warnings)
    assert any("Rust" in w for w in resume.fabrication_warnings)


def test_fabricated_position_removed(linkedin_data) -> None:
    fake = ResumePosition(company="Google", title="Staff Engineer", start_date="2020-01")
    real = _resume().positions[0]
    resume = check_fabrication(_resume(positions=[fake, real], skills=["React"]), linkedin_data)
    assert [p.company for p in resume.positions] == ["Acme Corp"]
    assert any("Google" in w for w in resume.fabrication_warnings)


def test_supplemental_tags_count_as_evidence(linkedin_data) -> None:
    merged = merge_additional_bullets(linkedin_data, {"Acme Corp": []})
    merged.supplemental_tags = ["Vite"]
    resume = check_fabrication(_resume(skills=["Vite"]), merged)
    assert resume.skills == ["Vite"]


def test_assemble_uses_source_identity(linkedin_data) -> None:
    resume = assemble_resume({
        "full_name": "Someone Else", "headline": "Lead", "summary": "s",
        "positions": [], "skills": ["React"],
        "education": [{"school": "Fake U"}],
    }, linkedin_data)
    assert resume.full_name == "Ada Lovelace"
    assert [e.school for e in resume.education] == ["State University"]
    assert "ada@example.com" in resume.contact_line


# ---------------------------------------------------------------------------
# Formatter
# ---------------------------------------------------------------------------

def _full_resume(linkedin_data) -> TailoredResume:
    return assemble_resume({
        "headline": "Frontend Lead for Acme",
        "summary": "Seasoned frontend lead.",
        "positions": [
            {"company": p.company, "title": p.title, "location": p.location,
             "start_date": p.start_date, "end_date": p.end_date,
             "bullets": [{"text": b} for b in p.bullets]}
            for p in linkedin_data.positions
        ],
        "skills": ["React", "JavaScript", "AEM"],
    }, linkedin_data)


def test_render_pdf_and_docx(linkedin_data, tmp_path: Path) -> None:
    import docx
    import pymupdf

    resume = _full_resume(linkedin_data)
    pdf = render_pdf(resume, tmp_path / "r.pdf")
    out_docx = render_docx(resume, tmp_path / "r.docx")
    assert pdf.stat().st_size > 0 and out_docx.stat().st_size > 0

    with pymupdf.open(pdf) as doc:
        text = "".join(page.get_text() for page in doc)
    assert "Ada Lovelace" in text
    for p in resume.positions:
        assert p.company in text
    assert "•" in text
    assert text.index("Summary") < text.index("Skills") < text.index("Experience") \
        < text.index("Education") < text.index("Certifications")

    d = docx.Document(str(out_docx))
    assert not d.tables
    doc_text = "\n".join(p.text for p in d.paragraphs)
    assert "Ada Lovelace" in doc_text and "Acme Corp" in doc_text
    assert "\tApr 2022 – Present" in doc_text  # right-aligned via tab stop


# ---------------------------------------------------------------------------
# Pipeline cache
# ---------------------------------------------------------------------------

async def test_pipeline_caches_and_forces(linkedin_data, tmp_path: Path) -> None:
    from app.resume_generator.pipeline import ResumePipeline

    lane = SearchLaneConfig(name="frontend_developer", target_roles=["Frontend Lead"],
                            target_field="web development")
    settings = AppSettings(
        profile=ProfileConfig(search_lanes={"frontend_developer": lane}),
        boards=BoardsConfig(), secrets=SecretsConfig(),
    )

    claude = AsyncMock()
    claude.generate_tailored_resume.return_value = ResumeGenResult(
        resume_json={"headline": "Lead", "summary": "s", "skills": ["React"],
                     "positions": [], "matched_requirements": ["React"]},
        model_used="test-model", input_tokens=10, output_tokens=5, cost_estimate=0.001,
    )

    async with Database(tmp_path / "jobs.db") as db:
        job_id = await db.insert_job(JobRecord(
            source="indeed", url="https://x/1", title="Frontend Lead", company="Acme",
            description="Need React", search_lane="frontend_developer",
        ))
        other_id = await db.insert_job(JobRecord(
            source="indeed", url="https://x/2", title="Marketing Manager",
            search_lane="marketing_manager",
        ))

        pipeline = ResumePipeline(
            db, claude, settings,
            linkedin_path=tmp_path / "linkedin_data.json",
            bullets_path=tmp_path / "no_bullets.json",
            output_dir=tmp_path / "out",
        )

        first = await pipeline.generate_for_job(job_id, "frontend_developer")
        assert not first.from_cache and first.gen_result is not None
        assert first.pdf_path.exists() and first.docx_path.exists()
        assert first.pdf_path.with_suffix(".json").exists()

        second = await pipeline.generate_for_job(job_id, "frontend_developer")
        assert second.from_cache and second.gen_result is None
        assert second.resume == first.resume
        assert claude.generate_tailored_resume.await_count == 1

        forced = await pipeline.generate_for_job(job_id, "frontend_developer", force=True)
        assert not forced.from_cache
        assert claude.generate_tailored_resume.await_count == 2

        doc = await db.get_generated_document(job_id, "frontend_developer", "resume")
        assert doc["id"] == forced.doc_id  # old row replaced, not duplicated
        cursor = await db.conn.execute("SELECT COUNT(*) FROM generated_documents")
        assert (await cursor.fetchone())[0] == 1

        with pytest.raises(ValueError, match="lane"):
            await pipeline.generate_for_job(other_id, "frontend_developer")
